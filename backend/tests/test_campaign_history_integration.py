import io
import openpyxl
import pytest
from httpx import AsyncClient

from app.models.campaign import Campaign
from app.services.campaign_workflow_service import CampaignWorkflowService, WorkflowStepKey, NextStepKey


from tests.test_creator_discovery import register


def _create_multi_sheet_excel() -> bytes:
    wb = openpyxl.Workbook()
    # Sheet 1: Campaign_Data
    ws_camp = wb.active
    ws_camp.title = "Campaign_Data"

    camp_headers = [
        "Campaign ID", "Campaign Name", "Brand", "Platform", "Start Date", "End Date",
        "Budget", "Actual Spend", "Revenue Generated", "Influencer", "Influencers Selected",
        "Campaign Stage", "Outreach Status", "Contract Status", "Views / Reach", "Status"
    ]
    ws_camp.append(camp_headers)

    camp_rows = [
        ["CMP-001", "Summer Hydration Wave", "AquaBlast", "YouTube", "2024-06-01", "2024-06-30", 50000, 48000, 144000, "AquaVlogger", 1, "Completed", "Accepted", "Signed", 250000, "Completed"],
        ["CMP-002", "Back to School Fuel", "SmartSnacks", "YouTube", "2024-08-01", "2024-08-31", 30000, 29000, 87000, "StudyWithMe", 1, "Completed", "Accepted", "Signed", 180000, "Completed"],
        ["CMP-003", "Fall Protein Blitz", "IronFit", "YouTube", "2024-09-01", "2024-09-30", 40000, 39500, 118500, "FitBeast", 1, "Completed", "Accepted", "Signed", 210000, "Completed"],
        ["CMP-004", "Holiday Glow 2024", "LumiSkin", "YouTube", "2024-11-01", "2024-11-30", 60000, 59000, 177000, "BeautyGlow", 1, "Completed", "Accepted", "Signed", 320000, "Completed"],
        ["CMP-005", "New Year Reset", "ZenLife", "YouTube", "2025-01-01", "2025-01-31", 35000, 34000, 102000, "MindfulLiving", 1, "Completed", "Accepted", "Signed", 190000, "Completed"],
        ["CMP-006", "Zero Sugar Refresh", "RefreshCo", "YouTube", "2025-02-01", "2025-02-28", 45000, 0, 0, "", 0, "Influencer Discovery", "", "", 0, "In Progress"],
    ]
    for r in camp_rows:
        ws_camp.append(r)

    # Sheet 2: Influencer_Discovery
    ws_disc = wb.create_sheet(title="Influencer_Discovery")
    disc_headers = [
        "Campaign ID", "Creator / Channel Name", "Channel Link", "Category / Niche",
        "Follower / Subscriber Count", "Avg Views", "Engagement Rate", "Audience Fit Score",
        "Match Score", "Predicted ROAS", "Discovery Decision", "Outreach Result"
    ]
    ws_disc.append(disc_headers)

    disc_rows = [
        ["CMP-006", "FitWithAlex", "https://youtube.com/@fitwithalex", "Fitness", 250000, 45000, 4.5, 92.0, 94.0, 3.4, "Recommended", "Not Contacted"],
        ["CMP-006", "HealthyEatsDaily", "https://youtube.com/@healthyeats", "Nutrition", 180000, 32000, 3.8, 88.0, 89.0, 3.1, "Recommended", "Not Contacted"],
        ["CMP-006", "HydrateLife", "https://youtube.com/@hydratelife", "Wellness", 310000, 55000, 5.1, 95.0, 96.0, 3.8, "Recommended", "Not Contacted"],
        ["CMP-006", "DailyVibeYoga", "https://youtube.com/@dailyvibeyoga", "Yoga", 120000, 22000, 4.2, 86.0, 87.0, 2.9, "Recommended", "Not Contacted"],
    ]
    for r in disc_rows:
        ws_disc.append(r)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_full_campaign_history_lifecycle(client: AsyncClient):
    headers = await register(client, "integration.history@test.com")
    excel_bytes = _create_multi_sheet_excel()

    # 1. Upload multi-sheet workbook
    upload_resp = await client.post(
        "/api/v1/assistant/imports",
        headers=headers,
        files={"files": ("Campaign_History.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert upload_resp.status_code == 200, upload_resp.text
    preview = upload_resp.json()

    assert preview["detected_campaigns"] == 6
    assert preview["completed_campaigns"] == 5
    assert preview["in_progress_campaigns"] == 1
    assert preview["needs_review"] == 0

    cands_by_name = {c["campaign_name"]: c for c in preview["campaigns"]}

    # Verify CMP-001 to CMP-005 are classified as COMPLETED
    for name in [
        "Summer Hydration Wave",
        "Back to School Fuel",
        "Fall Protein Blitz",
        "Holiday Glow 2024",
        "New Year Reset",
    ]:
        cand = cands_by_name[name]
        assert cand["classification"] == "COMPLETED", f"{name} should be COMPLETED"
        assert cand["current_stage"] == "COMPLETE"
        assert cand["continue_route"] is None

    # Verify CMP-006 (Zero Sugar Refresh)
    cmp006 = cands_by_name["Zero Sugar Refresh"]
    assert cmp006["classification"] == "IN_PROGRESS"
    assert cmp006["current_stage"] == "SHORTLIST"
    assert cmp006["workflow_state"] == "DISCOVERY_COMPLETED"
    assert cmp006["next_step_key"] == "SHORTLIST_INFLUENCERS"
    assert len(cmp006["creators"]) == 4

    creator_names = {cr["name"] for cr in cmp006["creators"]}
    assert "FitWithAlex" in creator_names
    assert "HealthyEatsDaily" in creator_names
    assert "HydrateLife" in creator_names
    assert "DailyVibeYoga" in creator_names

    alex = next(cr for cr in cmp006["creators"] if cr["name"] == "FitWithAlex")
    assert alex["followers"] == 250000
    assert alex["category"] == "Fitness"
    assert alex["discovery_decision"] == "Recommended"
    assert alex["match_score"] == 94.0

    # 2. Confirm import
    import_id = preview["import_id"]
    confirm_resp = await client.post(
        f"/api/v1/assistant/imports/{import_id}/confirm",
        headers=headers,
        json={"campaigns": []},
    )
    assert confirm_resp.status_code == 200, confirm_resp.text
    confirm_data = confirm_resp.json()
    assert len(confirm_data["campaign_ids"]) == 6

    # 3. List campaigns and check workflow state for each
    list_resp = await client.get("/api/v1/campaigns", headers=headers)
    assert list_resp.status_code == 200
    db_camps = {c["name"]: c for c in list_resp.json()}

    # Check completed campaigns
    summer_camp = db_camps["Summer Hydration Wave"]
    assert summer_camp["status"] == "completed"
    assert summer_camp["progress"] == 100

    wf_resp = await client.get(f"/api/v1/campaigns/{summer_camp['id']}/workflow", headers=headers)
    assert wf_resp.status_code == 200
    wf_data = wf_resp.json()
    assert wf_data["progress_percentage"] == 100
    assert wf_data["next_action"]["label"] == "Campaign Completed"
    assert all(s["status"] == "COMPLETED" for s in wf_data["steps"])

    # Check CMP-006 (Zero Sugar Refresh)
    cmp006_db = db_camps["Zero Sugar Refresh"]
    assert cmp006_db["status"] != "completed"
    wf006_resp = await client.get(f"/api/v1/campaigns/{cmp006_db['id']}/workflow", headers=headers)
    assert wf006_resp.status_code == 200
    wf006_data = wf006_resp.json()
    assert wf006_data["current_step"] == "SHORTLIST"
    assert wf006_data["next_step"] == "SHORTLIST_INFLUENCERS"
    assert wf006_data["next_action"]["label"] == "Review Influencers"
    assert wf006_data["next_action"]["tab"] == "influencers"
    assert wf006_data["discovered_count"] == 4
    assert wf006_data["shortlisted_count"] == 0

    # Discovery step must be COMPLETED
    discovery_step = next(s for s in wf006_data["steps"] if s["key"] == "DISCOVERY")
    assert discovery_step["status"] == "COMPLETED"

    # Shortlist step must be NEXT
    shortlist_step = next(s for s in wf006_data["steps"] if s["key"] == "SHORTLIST")
    assert shortlist_step["status"] == "NEXT"

    # 4. Check creators on CMP-006
    creators_resp = await client.get(f"/api/v1/campaigns/{cmp006_db['id']}/influencers", headers=headers)
    assert creators_resp.status_code == 200
    creators_data = creators_resp.json()["creators"]
    assert len(creators_data) == 4
    for cr in creators_data:
        assert cr["status"] == "DISCOVERED"
        assert cr["match_score"] is not None
        assert cr["creator"]["followers"] > 0
        assert cr["match_reasons"] is not None
        assert len(cr["match_reasons"]) > 0

    # 5. Test Assistant grounded Q&A
    # Question A: "Is Zero Sugar Refresh completed?"
    chat1 = await client.post(
        "/api/v1/assistant/chat",
        headers=headers,
        json={"message": "Is Zero Sugar Refresh completed?"},
    )
    assert chat1.status_code == 200
    reply1 = chat1.json()["reply"]
    assert "No." in reply1
    assert "Zero Sugar Refresh is not completed" in reply1
    assert "Influencer Discovery is complete" in reply1
    assert "4 creator" in reply1
    assert "No creator is recorded as selected yet" in reply1
    assert "Outreach and Contract have not started" in reply1

    # Question B: "Is Summer Hydration Wave completed?"
    chat2 = await client.post(
        "/api/v1/assistant/chat",
        headers=headers,
        json={"message": "Is Summer Hydration Wave completed?"},
    )
    assert chat2.status_code == 200
    reply2 = chat2.json()["reply"]
    assert "Yes." in reply2
    assert "Summer Hydration Wave is completed" in reply2

    # Question C: "Which campaigns are completed?"
    chat3 = await client.post(
        "/api/v1/assistant/chat",
        headers=headers,
        json={"message": "Which campaigns are completed?"},
    )
    assert chat3.status_code == 200
    reply3 = chat3.json()["reply"]
    assert "5 completed campaign(s)" in reply3

    # Question D: "Which campaigns still need work?"
    chat4 = await client.post(
        "/api/v1/assistant/chat",
        headers=headers,
        json={"message": "Which campaigns still need work?"},
    )
    assert chat4.status_code == 200
    reply4 = chat4.json()["reply"]
    assert "1 campaign(s) still need work" in reply4
    assert "Zero Sugar Refresh" in reply4


@pytest.mark.asyncio
async def test_multi_file_campaign_history_upload(client: AsyncClient):
    headers = await register(client, "multifile.history@test.com")

    # File 1: Campaigns.xlsx
    wb1 = openpyxl.Workbook()
    ws1 = wb1.active
    ws1.title = "Campaigns"
    ws1.append(["Campaign ID", "Campaign Name", "Brand", "Budget", "Actual Spend", "Revenue Generated", "Status"])
    ws1.append(["CMP-MULTI-01", "EcoBottle Launch", "GreenEarth", 25000, 24000, 75000, "Completed"])
    buf1 = io.BytesIO()
    wb1.save(buf1)

    # File 2: Influencers.xlsx
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "Influencers"
    ws2.append(["Campaign ID", "Influencer Name", "Channel Link", "Followers", "Category", "Discovery Decision", "Outreach Status"])
    ws2.append(["CMP-MULTI-01", "EcoWarrior", "https://youtube.com/@ecowarrior", 150000, "Sustainability", "Selected", "Accepted"])
    buf2 = io.BytesIO()
    wb2.save(buf2)

    # File 3: Performance.xlsx
    wb3 = openpyxl.Workbook()
    ws3 = wb3.active
    ws3.title = "Performance"
    ws3.append(["Campaign ID", "Influencer Name", "Views / Reach", "Revenue Generated", "Compensation", "Contract Status"])
    ws3.append(["CMP-MULTI-01", "EcoWarrior", 120000, 75000, 24000, "Signed"])
    buf3 = io.BytesIO()
    wb3.save(buf3)

    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    upload_resp = await client.post(
        "/api/v1/assistant/imports",
        headers=headers,
        files=[
            ("files", ("Campaigns.xlsx", buf1.getvalue(), mime)),
            ("files", ("Influencers.xlsx", buf2.getvalue(), mime)),
            ("files", ("Performance.xlsx", buf3.getvalue(), mime)),
        ],
    )
    assert upload_resp.status_code == 200, upload_resp.text
    preview = upload_resp.json()

    assert preview["detected_campaigns"] == 1
    cand = preview["campaigns"][0]
    assert cand["campaign_name"] == "EcoBottle Launch"
    assert cand["classification"] == "COMPLETED"
    assert len(cand["creators"]) == 1
    assert cand["creators"][0]["name"] == "EcoWarrior"
    assert cand["creators"][0]["followers"] == 150000

