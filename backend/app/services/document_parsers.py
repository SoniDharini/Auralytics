"""Deterministic parsers for campaign-history files. Never invents rows."""

from __future__ import annotations

import csv
import io
import json
import re
from typing import Any, Dict, List, Optional, Tuple

ALLOWED_EXTENSIONS = {".csv", ".xlsx", ".xls", ".json", ".pdf", ".docx"}
ALLOWED_MIME = {
    "text/csv",
    "application/csv",
    "text/plain",
    "application/json",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel",
    "application/excel",
    "application/x-excel",
    "application/x-msexcel",
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/octet-stream",
    "application/zip",
}
MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_FILES = 12

_INJECTION_MARKERS = (
    "ignore all previous instructions",
    "ignore your instructions",
    "reveal api key",
    "give me api keys",
    "system prompt",
)


class ParsedTable:
    def __init__(self, filename: str, file_type: str, rows: List[Dict[str, Any]], text: str = "", sheet: str = ""):
        self.filename = filename
        self.file_type = file_type
        self.rows = rows
        self.text = text
        self.sheet = sheet


def classify_source_type(filename: str) -> str:
    name = filename.lower()
    if "contract" in name:
        return "CONTRACT"
    if "outreach" in name or "negotiat" in name:
        return "OUTREACH"
    if "performance" in name or "analytics" in name or "report" in name:
        return "PERFORMANCE"
    if "creator" in name or "influencer" in name:
        return "CREATOR"
    if "optim" in name:
        return "OPTIMIZATION"
    return "CAMPAIGN"


def validate_upload(filename: str, content_type: Optional[str], data: bytes) -> str:
    if not filename or "." not in filename:
        raise ValueError("File name is missing an extension.")
    ext = "." + filename.rsplit(".", 1)[-1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file type '{ext}'. Use XLSX, CSV, JSON, PDF, or DOCX.")
    if content_type:
        mime = content_type.split(";")[0].strip().lower()
        if mime and mime not in ALLOWED_MIME and ext not in {".xlsx", ".xls", ".csv", ".docx", ".pdf", ".json"}:
            raise ValueError("Unsupported MIME type.")
    if not data:
        raise ValueError("The uploaded file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("File exceeds the 8 MB limit.")
    if data[:2] == b"MZ" or b"\x00javascript" in data[:200].lower():
        raise ValueError("Executable or script content is not allowed.")
    return ext.lstrip(".")


def parse_file(filename: str, data: bytes, file_type: str) -> List[ParsedTable]:
    kind = file_type.lower()
    if kind == "csv":
        return [_parse_csv(filename, data)]
    if kind == "json":
        return [_parse_json(filename, data)]
    if kind in ("xlsx", "xls"):
        return _parse_xlsx(filename, data)
    if kind == "pdf":
        return [_parse_pdf(filename, data)]
    if kind == "docx":
        return [_parse_docx(filename, data)]
    raise ValueError(f"Unsupported file type '{file_type}'.")


def _parse_csv(filename: str, data: bytes) -> ParsedTable:
    text = _decode_text(data)
    reader = csv.DictReader(io.StringIO(text))
    rows = [_clean_row(row) for row in reader if any(str(v).strip() for v in row.values() if v is not None)]
    return ParsedTable(filename, "csv", rows, text=text)


def _parse_json(filename: str, data: bytes) -> ParsedTable:
    try:
        payload = json.loads(_decode_text(data))
    except json.JSONDecodeError as exc:
        raise ValueError(f"{filename} is not valid JSON.") from exc
    rows = _json_to_rows(payload)
    return ParsedTable(filename, "json", rows, text=_decode_text(data)[:20000])


def _json_to_rows(payload: Any) -> List[Dict[str, Any]]:
    if isinstance(payload, list):
        out: List[Dict[str, Any]] = []
        for item in payload:
            if isinstance(item, dict):
                out.extend(_flatten_campaign_object(item))
        return out
    if isinstance(payload, dict):
        if isinstance(payload.get("campaigns"), list):
            rows: List[Dict[str, Any]] = []
            for item in payload["campaigns"]:
                if isinstance(item, dict):
                    rows.extend(_flatten_campaign_object(item))
            return rows
        return _flatten_campaign_object(payload)
    raise ValueError("JSON must be an object or an array of campaign objects.")


def _flatten_campaign_object(item: Dict[str, Any]) -> List[Dict[str, Any]]:
    campaign = item.get("campaign") if isinstance(item.get("campaign"), dict) else item
    base = {k: v for k, v in campaign.items() if not isinstance(v, (list, dict))}
    if item.get("campaign_name") and "campaign_name" not in base:
        base["campaign_name"] = item.get("campaign_name")
    creators = item.get("creators") or item.get("influencers") or []
    outreach = item.get("outreach_records") or item.get("outreach") or []
    contracts = item.get("contracts") or []
    performance = item.get("performance_records") or item.get("performance") or []
    content = item.get("content_records") or item.get("content") or []
    if not any([creators, outreach, contracts, performance, content]):
        return [base] if base else []
    rows: List[Dict[str, Any]] = []
    max_len = max(1, len(creators), len(outreach), len(contracts), len(performance), len(content))
    for i in range(max_len):
        row = dict(base)
        if i < len(creators) and isinstance(creators[i], dict):
            row.update({f"creator_{k}" if k in base else k: v for k, v in creators[i].items()})
            if creators[i].get("name") and "creator_name" not in row:
                row["creator_name"] = creators[i].get("name")
        if i < len(outreach) and isinstance(outreach[i], dict):
            row.update(outreach[i])
        if i < len(contracts) and isinstance(contracts[i], dict):
            row.update({f"contract_{k}" if k in row else k: v for k, v in contracts[i].items()})
        if i < len(performance) and isinstance(performance[i], dict):
            row.update(performance[i])
        if i < len(content) and isinstance(content[i], dict):
            row.update(content[i])
        rows.append(row)
    return rows


def _parse_xlsx(filename: str, data: bytes) -> List[ParsedTable]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return _parse_xlsx_stdlib(filename, data)
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        try:
            return _parse_xlsx_stdlib(filename, data)
        except Exception:
            raise ValueError(f"{filename} could not be read as an Excel workbook.") from exc
    tables: List[ParsedTable] = []
    for sheet in wb.worksheets:
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            continue
        headers = [str(h).strip() if h is not None else "" for h in header_row]
        if not any(headers):
            continue
        rows: List[Dict[str, Any]] = []
        for raw in rows_iter:
            row = {headers[i]: raw[i] for i in range(min(len(headers), len(raw))) if headers[i]}
            cleaned = _clean_row(row)
            if any(str(v).strip() for v in cleaned.values() if v is not None):
                rows.append(cleaned)
        if rows:
            tables.append(ParsedTable(filename, "xlsx", rows, sheet=sheet.title or ""))
    if not tables:
        raise ValueError(f"{filename} has no readable table rows.")
    return tables


def _parse_xlsx_stdlib(filename: str, data: bytes) -> List[ParsedTable]:
    """Parse .xlsx via zip/XML when openpyxl is unavailable. .xls is not supported here."""
    import zipfile
    from xml.etree import ElementTree as ET

    if not data.startswith(b"PK"):
        raise ValueError(
            f"{filename} looks like a legacy .xls file. Save it as .xlsx and upload again."
        )
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            shared: List[str] = []
            if "xl/sharedStrings.xml" in zf.namelist():
                root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
                for si in root.findall("m:si", ns):
                    shared.append("".join(node.text or "" for node in si.iter() if node.text))
            sheets: List[Tuple[str, str]] = []
            if "xl/workbook.xml" in zf.namelist() and "xl/_rels/workbook.xml.rels" in zf.namelist():
                rels_root = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
                rel_ns = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}
                rels = {
                    rel.attrib.get("Id"): rel.attrib.get("Target")
                    for rel in rels_root.findall("r:Relationship", rel_ns)
                }
                wb_root = ET.fromstring(zf.read("xl/workbook.xml"))
                r_ns = {"r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
                for sheet in wb_root.findall("m:sheets/m:sheet", ns):
                    rid = sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}Id") or sheet.attrib.get(
                        f"{{{r_ns['r']}}}id"
                    )
                    target = rels.get(rid or "", "")
                    path = target if target.startswith("xl/") else f"xl/{target.lstrip('/')}"
                    sheets.append((sheet.attrib.get("name") or "Sheet", path))
            if not sheets:
                sheets = [
                    ("Sheet1", name)
                    for name in zf.namelist()
                    if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
                ]
            tables: List[ParsedTable] = []
            for title, path in sheets:
                if path not in zf.namelist():
                    continue
                rows = _xlsx_sheet_rows(ET.fromstring(zf.read(path)), shared, ns)
                if rows:
                    tables.append(ParsedTable(filename, "xlsx", rows, sheet=title))
    except zipfile.BadZipFile as exc:
        raise ValueError(f"{filename} could not be read as an Excel workbook.") from exc
    if not tables:
        raise ValueError(f"{filename} has no readable table rows.")
    return tables


def _xlsx_sheet_rows(sheet_root: Any, shared: List[str], ns: Dict[str, str]) -> List[Dict[str, Any]]:
    def col_index(cell_ref: str) -> int:
        letters = "".join(ch for ch in cell_ref if ch.isalpha())
        idx = 0
        for ch in letters:
            idx = idx * 26 + (ord(ch.upper()) - 64)
        return max(idx - 1, 0)

    matrix: List[List[Any]] = []
    for row_el in sheet_root.findall("m:sheetData/m:row", ns):
        values: Dict[int, Any] = {}
        max_idx = -1
        for cell in row_el.findall("m:c", ns):
            ref = cell.attrib.get("r") or ""
            idx = col_index(ref) if ref else (max_idx + 1)
            max_idx = max(max_idx, idx)
            cell_type = cell.attrib.get("t")
            value_el = cell.find("m:v", ns)
            is_el = cell.find("m:is", ns)
            if is_el is not None:
                values[idx] = "".join(node.text or "" for node in is_el.iter() if node.text)
            elif value_el is not None and value_el.text is not None:
                raw = value_el.text
                if cell_type == "s":
                    try:
                        values[idx] = shared[int(raw)]
                    except (ValueError, IndexError):
                        values[idx] = raw
                else:
                    values[idx] = raw
        if max_idx < 0:
            continue
        line = [values.get(i) for i in range(max_idx + 1)]
        matrix.append(line)
    if not matrix:
        return []
    headers = [str(h).strip() if h is not None else "" for h in matrix[0]]
    rows: List[Dict[str, Any]] = []
    for raw in matrix[1:]:
        row = {headers[i]: raw[i] if i < len(raw) else None for i in range(len(headers)) if headers[i]}
        cleaned = _clean_row(row)
        if any(str(v).strip() for v in cleaned.values() if v is not None):
            rows.append(cleaned)
    return rows


def _parse_pdf(filename: str, data: bytes) -> ParsedTable:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ValueError("PDF support is not installed on the server.") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ValueError(f"{filename} could not be read as a PDF.") from exc
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            continue
    text = "\n".join(pages).strip()
    if not text:
        raise ValueError(f"{filename} contains no extractable text.")
    return ParsedTable(filename, "pdf", _text_to_rows(text), text=text)


def _parse_docx(filename: str, data: bytes) -> ParsedTable:
    try:
        from docx import Document
    except ImportError as exc:
        raise ValueError("DOCX support is not installed on the server.") from exc
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ValueError(f"{filename} could not be read as a Word document.") from exc
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        headers = [c.text.strip() for c in table.rows[0].cells] if table.rows else []
        for row in table.rows[1:]:
            values = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(f"{headers[i]}: {values[i]}" if i < len(headers) else values[i] for i in range(len(values))))
    text = "\n".join(parts)
    if not text.strip():
        raise ValueError(f"{filename} contains no extractable text.")
    return ParsedTable(filename, "docx", _text_to_rows(text), text=text)


def _text_to_rows(text: str) -> List[Dict[str, Any]]:
    chunks = re.split(r"\n(?=campaign\s*name\s*:|campaign\s*:)", text, flags=re.IGNORECASE)
    rows: List[Dict[str, Any]] = []
    for chunk in chunks:
        row: Dict[str, Any] = {}
        for line in chunk.splitlines():
            if ":" not in line:
                continue
            key, value = line.split(":", 1)
            key = key.strip()
            value = value.strip()
            if key and value:
                row[key] = value
        if row:
            rows.append(_clean_row(row))
    if not rows and text.strip():
        rows.append({"unstructured_text": text[:8000]})
    return rows


def sanitize_untrusted_text(text: str) -> str:
    """Keep document text as data. Neutralize common prompt-injection phrases for LLM context."""
    lowered = text
    for marker in _INJECTION_MARKERS:
        pattern = re.compile(re.escape(marker), re.IGNORECASE)
        lowered = pattern.sub("[untrusted document text omitted]", lowered)
    return lowered


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ValueError("The file encoding could not be decoded.")


def _clean_row(row: Dict[str, Any]) -> Dict[str, Any]:
    cleaned: Dict[str, Any] = {}
    for key, value in row.items():
        if key is None:
            continue
        name = str(key).strip()
        if not name:
            continue
        if value is None:
            continue
        if isinstance(value, str):
            value = value.strip()
            if not value or value.lower() in {"n/a", "na", "none", "null", "unknown", "-"}:
                continue
        cleaned[name] = value
    return cleaned
